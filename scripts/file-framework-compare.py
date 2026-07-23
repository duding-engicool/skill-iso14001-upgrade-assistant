#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ISO 14001:2026 文件框架差距比对脚本

功能：解析用户上传的现有体系文件框架（.md/.xlsx/.docx/.txt），
      与内置的 2026 要求清单做关键词映射，输出差距矩阵（markdown）。

用法：
  python file-framework-compare.py <框架文件> [-o 输出.md]

说明：
  - .md/.txt 用标准库即可解析（提取标题/列表/段落条目）。
  - .xlsx 需 openpyxl：pip install openpyxl
  - .docx 需 python-docx：pip install python-docx
  - 比对结果为辅助定位，存在伪匹配/漏匹配可能，输出后需人工逐项确认。
  - 基于 ISO 14001:2026 FDIS 草案，正式版发布后复核。
"""

import argparse
import os
import re
import sys

# 内置 2026 要求清单：条款名 → 关键词（命中任一即视为"有对应"）
REQUIREMENTS = [
    ("4.1 理解组织及其环境", ["组织环境", "context", "内外部", "内外部问题", "气候", "climate"]),
    ("4.2 相关方需求与期望", ["相关方", "interested", "需求", "期望"]),
    ("4.3 确定范围", ["范围", "scope"]),
    ("5.2 环境方针", ["方针", "policy"]),
    ("6.1.2 环境因素", ["环境因素", "环境因子", "aspect", "重要环境因素"]),
    ("6.1.3 合规义务", ["合规", "compliance", "法律法规", "义务"]),
    ("6.1.4 风险和机遇", ["风险", "机遇", "risk", "opportunity", "风险机遇"]),
    ("6.2 环境目标及实现策划", ["目标", "objective", "指标"]),
    ("7 支持(能力/意识/沟通/文件化)", ["能力", "意识", "沟通", "文件化", "培训", "competence"]),
    ("8.1 运行策划和控制", ["运行", "设计", "采购", "外包", "外部供方", "供方", "生命周期", "供应商"]),
    ("8.2 应急准备和响应", ["应急", "emergency", "预案"]),
    ("9.1 监视测量分析评价", ["监视", "测量", "monitor"]),
    ("9.2 内部审核", ["内审", "内部审核", "audit"]),
    ("9.3 管理评审", ["管理评审", "management review"]),
    ("10 改进(持续改进/纠正措施)", ["改进", "纠正", "corrective", "持续优化"]),
]


def _extract_text(path):
    """从 .md/.txt 提取条目：去标题/列表符号，逐行。"""
    entries = []
    with open(path, encoding="utf-8", errors="ignore") as f:
        for line in f:
            s = line.strip()
            if not s:
                continue
            s = re.sub(r"^[#\-\*\+\d+\.\s\[\]\(\]\)]+", "", s).strip()
            if len(s) >= 2:
                entries.append(s)
    return entries


def _extract_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        return None, "需要 openpyxl 才能解析 .xlsx，请运行：pip install openpyxl（或改用 .md/.txt）"
    wb = openpyxl.load_workbook(path, data_only=True)
    entries = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    s = str(cell).strip()
                    if len(s) >= 2:
                        entries.append(s)
    return entries, None


def _extract_docx(path):
    try:
        import docx
    except ImportError:
        return None, "需要 python-docx 才能解析 .docx，请运行：pip install python-docx（或改用 .md/.txt）"
    document = docx.Document(path)
    entries = []
    for para in document.paragraphs:
        s = para.text.strip()
        if len(s) >= 2:
            entries.append(s)
    return entries, None


def extract_entries(path):
    """按扩展名分发解析，返回 (entries, error)。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".md", ".txt"):
        return _extract_text(path), None
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext == ".docx":
        return _extract_docx(path)
    # 未知扩展名，尝试按文本读
    return _extract_text(path), None


def compare(entries):
    """对每个要求条款，扫描用户条目关键词命中。"""
    results = []
    for name, kws in REQUIREMENTS:
        matched = [e for e in entries if any(kw.lower() in e.lower() for kw in kws)]
        results.append((name, bool(matched), matched[:3]))
    return results


def build_markdown(entries, results):
    lines = [
        "# ISO 14001:2026 文件框架差距矩阵",
        "",
        "> ⚠️ 基于 FDIS 草案自动比对，结果需人工确认（可能存在伪匹配/漏匹配）。",
        f"> 解析到 **{len(entries)}** 个框架条目。",
        "",
        "| 2026 要求条款 | 对应状态 | 匹配到的框架条目(示例) |",
        "|---|---|---|",
    ]
    for name, ok, samples in results:
        status = "✅ 有对应" if ok else "❌ 缺失"
        sample_txt = "；".join(samples) if samples else "—"
        lines.append(f"| {name} | {status} | {sample_txt} |")

    missing = [name for name, ok, _ in results if not ok]
    lines.append("")
    lines.append("## 缺失项（升版重点补）")
    lines.append("")
    if missing:
        for m in missing:
            lines.append(f"- {m}")
    else:
        lines.append("无自动识别缺失项（仍需人工确认内容是否真正达到 2026 要求，警惕伪匹配）。")
    lines.append("")
    lines.append("---")
    lines.append("下一步建议：对 ✅ 项人工核实内容深度；对 ❌ 项按《体系文件修订映射方法论》新建/修订对应文件。")
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description="ISO 14001:2026 文件框架差距比对")
    parser.add_argument("framework", help="用户上传的体系文件框架(.md/.xlsx/.docx/.txt)")
    parser.add_argument("-o", "--out", help="输出 markdown 路径（默认打印到控制台）")
    args = parser.parse_args()

    if not os.path.exists(args.framework):
        print(f"错误：文件不存在 {args.framework}")
        sys.exit(1)

    entries, err = extract_entries(args.framework)
    if err:
        print(err)
        sys.exit(1)
    if not entries:
        print("未从文件中解析到任何条目，请检查文件格式或内容。")
        sys.exit(1)

    results = compare(entries)
    md = build_markdown(entries, results)

    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(md)
        print(f"已写入：{args.out}")
    else:
        print(md)


if __name__ == "__main__":
    main()
